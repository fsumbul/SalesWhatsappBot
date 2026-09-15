"""Private, tenant-scoped CSV/XLSX imports for conversational outreach.

This module deliberately owns only import storage and parsing.  It does not
create an outbound batch or cross the Meta boundary; the existing outreach
workflow/outbox remains the only sender of messages.
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from enum import StrEnum
from functools import lru_cache
from threading import Lock
from typing import Any, Protocol, cast
from uuid import UUID, uuid4

import phonenumbers
from sqlalchemy import DateTime, ForeignKey, Index, Integer, String, Text, UniqueConstraint, select
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.dialects.postgresql import UUID as PgUUID
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import Mapped, mapped_column

from src.core.config import get_settings
from src.core.db import Base
from src.core.mixins import TenantScoped, Timestamped, UUIDPrimaryKey

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000
IMPORT_RETENTION_DAYS = 30
_MAX_COLUMNS = 200
_MAX_UNCOMPRESSED_XLSX_BYTES = 50 * 1024 * 1024
_MAX_PHONE_CELL_CHARS = 64
_IMPORT_OBJECT_PREFIX = "campaign-imports/"
# This marker lets the application fail closed if an operator accidentally
# points the feature at an existing shared MinIO bucket. Lifecycle/policy
# changes are only safe for a bucket this feature owns.
_BUCKET_OWNERSHIP_MARKER = ".leadpulse-campaign-import-bucket-v1"
_BUCKET_OWNERSHIP_BODY = b"leadpulse private campaign import bucket v1\n"
CONSENT_SOURCES = frozenset(
    {"web_form", "written_record", "customer_relationship", "other"}
)


def campaign_imports_enabled_for(tenant_id: UUID, manager_id: UUID) -> bool:
    """Return whether this concrete canary permits one manager's file campaign.

    Empty allowlists never mean "all tenants". This is intentionally not a
    reusable feature-flag mechanism: the first real-message rollout must name
    the single tenant and sales manager that were sandbox-validated.
    """

    settings = get_settings()
    tenants = settings.campaign_imports_canary_tenant_ids_list
    managers = settings.campaign_imports_canary_manager_ids_list
    return bool(
        settings.campaign_imports_enabled
        # Keep the runtime guard as strict as production preflight. A
        # development/staging typo must not widen a deliberately one-manager
        # real-message canary merely because production validation is absent.
        and len(tenants) == 1
        and len(managers) == 1
        and tenant_id in tenants
        and manager_id in managers
    )


def campaign_import_recipient_limit() -> int:
    """Keep the canary ceiling below the format/parser hard limit."""

    return min(MAX_IMPORT_ROWS, get_settings().campaign_imports_max_recipients)


class CampaignImportStatus(StrEnum):
    QUEUED = "queued"
    PARSING = "parsing"
    AWAITING_MAPPING = "awaiting_mapping"
    READY = "ready"
    FAILED = "failed"


class CampaignImportRowStatus(StrEnum):
    ELIGIBLE = "eligible"
    INVALID = "invalid"
    DUPLICATE = "duplicate"
    BLOCKED = "blocked"


class CampaignImportError(ValueError):
    """Base error suitable for a route to translate to a safe 4xx response."""


class CampaignImportValidationError(CampaignImportError):
    """The uploaded artifact or selected mapping is not acceptable."""


class CampaignImportConflictError(CampaignImportError):
    """A client operation id is already bound to different upload input."""


class CampaignImportStorageError(RuntimeError):
    """Object storage is unavailable or not configured."""


class CampaignImport(Base, UUIDPrimaryKey, TenantScoped, Timestamped):
    __tablename__ = "chat_campaign_imports"
    __table_args__ = (
        UniqueConstraint(
            "workflow_id", "client_operation_id", name="uq_campaign_import_workflow_operation"
        ),
        Index("ix_campaign_imports_tenant_expires", "tenant_id", "expires_at"),
    )

    workflow_id: Mapped[UUID] = mapped_column(
        PgUUID(as_uuid=True), ForeignKey("chat_workflows.id", ondelete="CASCADE"), nullable=False
    )
    uploader_id: Mapped[UUID] = mapped_column(
        PgUUID(as_uuid=True), ForeignKey("users.id", ondelete="CASCADE"), nullable=False
    )
    client_operation_id: Mapped[UUID] = mapped_column(PgUUID(as_uuid=True), nullable=False)
    object_key: Mapped[str] = mapped_column(String(512), nullable=False, unique=True)
    original_filename: Mapped[str] = mapped_column(String(255), nullable=False)
    sha256: Mapped[str] = mapped_column(String(64), nullable=False)
    mime_type: Mapped[str] = mapped_column(String(120), nullable=False)
    status: Mapped[str] = mapped_column(
        String(32), nullable=False, default=CampaignImportStatus.QUEUED.value, index=True
    )
    # A worker must present this lease token before it can persist parsing
    # output. It prevents a stale parser from overwriting a newer recovery
    # attempt after the beat lease expires.
    parse_token: Mapped[UUID | None] = mapped_column(PgUUID(as_uuid=True), nullable=True)
    country_code: Mapped[str] = mapped_column(String(2), nullable=False)
    phone_column: Mapped[str | None] = mapped_column(String(255), nullable=True)
    columns: Mapped[list[str]] = mapped_column(JSONB, nullable=False, default=list)
    total_rows: Mapped[int] = mapped_column(Integer, nullable=False, default=0)
    eligible_count: Mapped[int] = mapped_column(Integer, nullable=False, default=0)
    invalid_count: Mapped[int] = mapped_column(Integer, nullable=False, default=0)
    duplicate_count: Mapped[int] = mapped_column(Integer, nullable=False, default=0)
    blocked_count: Mapped[int] = mapped_column(Integer, nullable=False, default=0)
    summary: Mapped[dict[str, Any]] = mapped_column(JSONB, nullable=False, default=dict)
    consent_source: Mapped[str | None] = mapped_column(String(120), nullable=True)
    consent_note: Mapped[str | None] = mapped_column(Text, nullable=True)
    failure_reason: Mapped[str | None] = mapped_column(Text, nullable=True)
    expires_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), nullable=False)


class CampaignImportRow(Base, UUIDPrimaryKey, TenantScoped, Timestamped):
    __tablename__ = "chat_campaign_import_rows"
    __table_args__ = (
        UniqueConstraint("campaign_import_id", "row_number", name="uq_campaign_import_row_number"),
        Index("ix_campaign_import_rows_import_status", "campaign_import_id", "status"),
    )

    campaign_import_id: Mapped[UUID] = mapped_column(
        PgUUID(as_uuid=True),
        ForeignKey("chat_campaign_imports.id", ondelete="CASCADE"),
        nullable=False,
        index=True,
    )
    row_number: Mapped[int] = mapped_column(Integer, nullable=False)
    phone_e164: Mapped[str | None] = mapped_column(String(32), nullable=True)
    masked_phone: Mapped[str | None] = mapped_column(String(32), nullable=True)
    status: Mapped[str] = mapped_column(String(20), nullable=False, index=True)
    reason: Mapped[str | None] = mapped_column(Text, nullable=True)


class CampaignImportStorage(Protocol):
    async def put_bytes(self, object_key: str, data: bytes, mime_type: str) -> None: ...

    async def get_bytes(self, object_key: str) -> bytes: ...

    async def delete(self, object_key: str) -> None: ...


class MinioCampaignImportStorage:
    """Small async facade over the synchronous MinIO client.

    The bucket remains private: this module never creates a presigned browser
    URL.  The web BFF/API streams the upload and workers use credentials only
    on trusted hosts.
    """

    def __init__(
        self,
        *,
        endpoint: str,
        access_key: str,
        secret_key: str,
        bucket: str,
        secure: bool,
    ) -> None:
        if not endpoint or not access_key or not secret_key or not bucket:
            raise CampaignImportStorageError("Private import storage is not configured.")
        self.endpoint = endpoint
        self.access_key = access_key
        self.secret_key = secret_key
        self.bucket = bucket
        self.secure = secure
        self._bucket_ready = False
        self._bucket_lock = Lock()

    def _client(self) -> Any:
        try:
            from minio import Minio
        except ModuleNotFoundError as exc:  # pragma: no cover - deployment dependency
            raise CampaignImportStorageError("MinIO client dependency is unavailable.") from exc
        return Minio(
            self.endpoint,
            access_key=self.access_key,
            secret_key=self.secret_key,
            secure=self.secure,
        )

    def _ensure_bucket(self) -> None:
        if self._bucket_ready:
            return
        # Multiple uploads can enter asyncio.to_thread at application start.
        # Serialize the one-time create so a harmless first-use race does not
        # turn one manager's upload into a storage error.
        with self._bucket_lock:
            if self._bucket_ready:
                return
            try:
                from minio.commonconfig import Filter
                from minio.lifecycleconfig import Expiration, LifecycleConfig, Rule

                client = self._client()
                created = not client.bucket_exists(self.bucket)
                if created:
                    client.make_bucket(self.bucket)
                try:
                    client.stat_object(self.bucket, _BUCKET_OWNERSHIP_MARKER)
                except Exception as exc:
                    if getattr(exc, "code", None) not in {"NoSuchKey", "NoSuchObject"}:
                        raise
                    # Claiming an existing non-empty bucket would make its
                    # policy/lifecycle someone else's data-retention rule.
                    if not created and any(client.list_objects(self.bucket, recursive=True)):
                        raise CampaignImportStorageError(
                            "MinIO içe aktarma bucket'ı yalnız bu özellik için ayrılmış olmalı."
                        ) from exc
                    client.put_object(
                        self.bucket,
                        _BUCKET_OWNERSHIP_MARKER,
                        io.BytesIO(_BUCKET_OWNERSHIP_BODY),
                        length=len(_BUCKET_OWNERSHIP_BODY),
                        content_type="text/plain",
                    )
                try:
                    # MinIO/S3's default after removing a bucket policy is
                    # private. This actively removes a stale public policy;
                    # it never grants anonymous read/write access.
                    client.delete_bucket_policy(self.bucket)
                except Exception as exc:
                    if getattr(exc, "code", None) not in {"NoSuchBucketPolicy", "NoSuchPolicy"}:
                        raise
                # This is a dedicated private import bucket. Enforce object
                # expiry in storage as well as the DB cleanup job, so a crash
                # after MinIO write but before DB commit cannot retain a raw
                # phone file indefinitely.
                client.set_bucket_lifecycle(
                    self.bucket,
                    LifecycleConfig(
                        [
                            Rule(
                                status="Enabled",
                                rule_id="campaign-import-retention-30-days",
                                rule_filter=Filter(prefix=_IMPORT_OBJECT_PREFIX),
                                expiration=Expiration(days=IMPORT_RETENTION_DAYS),
                            )
                        ]
                    ),
                )
            except CampaignImportStorageError:
                raise
            except Exception as exc:  # pragma: no cover - MinIO configuration boundary
                raise CampaignImportStorageError(
                    "Private import storage retention policy could not be verified."
                ) from exc
            self._bucket_ready = True

    async def put_bytes(self, object_key: str, data: bytes, mime_type: str) -> None:
        if not object_key.startswith(_IMPORT_OBJECT_PREFIX):
            raise CampaignImportStorageError("Geçersiz özel içe aktarma nesne anahtarı.")

        def _put() -> None:
            self._ensure_bucket()
            self._client().put_object(
                self.bucket,
                object_key,
                io.BytesIO(data),
                length=len(data),
                content_type=mime_type,
            )

        try:
            await asyncio.to_thread(_put)
        except CampaignImportStorageError:
            raise
        except Exception as exc:  # pragma: no cover - depends on object storage availability
            raise CampaignImportStorageError("Private import storage could not save the file.") from exc

    async def get_bytes(self, object_key: str) -> bytes:
        if not object_key.startswith(_IMPORT_OBJECT_PREFIX):
            raise CampaignImportStorageError("Geçersiz özel içe aktarma nesne anahtarı.")

        def _get() -> bytes:
            self._ensure_bucket()
            response = self._client().get_object(self.bucket, object_key)
            try:
                return bytes(response.read())
            finally:
                response.close()
                response.release_conn()

        try:
            return await asyncio.to_thread(_get)
        except CampaignImportStorageError:
            raise
        except Exception as exc:  # pragma: no cover - depends on object storage availability
            raise CampaignImportStorageError("Private import storage could not read the file.") from exc

    async def delete(self, object_key: str) -> None:
        if not object_key.startswith(_IMPORT_OBJECT_PREFIX):
            raise CampaignImportStorageError("Geçersiz özel içe aktarma nesne anahtarı.")

        def _delete() -> None:
            self._ensure_bucket()
            try:
                self._client().remove_object(self.bucket, object_key)
            except Exception as exc:
                # Retention may resume after a process died between object
                # deletion and the DB commit. A missing object is therefore
                # already-cleaned, not a reason to retain its DB rows.
                if getattr(exc, "code", None) not in {"NoSuchKey", "NoSuchObject"}:
                    raise

        try:
            await asyncio.to_thread(_delete)
        except CampaignImportStorageError:
            raise
        except Exception as exc:  # pragma: no cover - depends on object storage availability
            raise CampaignImportStorageError("Private import storage could not delete the file.") from exc


@lru_cache
def get_campaign_import_storage() -> CampaignImportStorage:
    settings = get_settings()
    return MinioCampaignImportStorage(
        endpoint=settings.minio_endpoint,
        access_key=settings.minio_access_key,
        secret_key=settings.minio_secret_key,
        bucket=settings.minio_bucket,
        secure=settings.minio_secure,
    )


@dataclass(frozen=True)
class ParsedImport:
    columns: list[str]
    phone_column: str | None
    phone_values: list[tuple[int, str]]


_CSV_MIMES = {"text/csv", "application/csv", "text/plain", "application/vnd.ms-excel"}
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PHONE_HEADER_KEYS = {
    "phone",
    "phonenumber",
    "telephone",
    "telefon",
    "telefonnumarasi",
    "telefonnumarası",
    "gsm",
    "mobile",
    "mobilephone",
    "ceptelefonu",
    "whatsapp",
    "whatsappnumber",
    "numara",
}


def safe_filename(filename: str) -> str:
    """Keep only a basename and printable characters for metadata display."""

    name = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    name = "".join(char for char in name if char.isprintable())
    if not name or len(name) > 255:
        raise CampaignImportValidationError("Dosya adı geçerli değil.")
    return name


def _normalise_country_code(country_code: str) -> str:
    country = country_code.strip().upper()
    if not re.fullmatch(r"[A-Z]{2}", country) or not phonenumbers.country_code_for_region(country):
        raise CampaignImportValidationError("Geçerli bir ülke kodu seçin.")
    return country


def _zip_members(data: bytes) -> set[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            total_uncompressed = sum(info.file_size for info in infos)
            if total_uncompressed > _MAX_UNCOMPRESSED_XLSX_BYTES:
                raise CampaignImportValidationError("Excel dosyası güvenli ayrıştırma sınırını aşıyor.")
            if any(
                info.compress_size and info.file_size / info.compress_size > 100
                for info in infos
            ):
                raise CampaignImportValidationError("Excel dosyası güvenli ayrıştırma sınırını aşıyor.")
            return {info.filename.lower() for info in infos}
    except CampaignImportValidationError:
        raise
    except zipfile.BadZipFile as exc:
        raise CampaignImportValidationError("Geçerli bir XLSX dosyası yükleyin.") from exc


def validate_upload(filename: str, claimed_mime: str | None, data: bytes) -> tuple[str, str]:
    """Validate a private CSV/XLSX upload before it reaches object storage.

    CSV has no dependable magic signature, so it is checked again by the
    bounded parser. XLSX requires an OOXML zip with no macro or encryption
    payload. The return value is ``(safe_filename, canonical_mime)``.
    """

    name = safe_filename(filename)
    if not data or len(data) > MAX_IMPORT_BYTES:
        raise CampaignImportValidationError("Dosya en fazla 10 MiB olabilir.")
    suffix = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    mime = (claimed_mime or "").split(";", 1)[0].strip().lower()
    if suffix == "csv":
        if mime and mime not in _CSV_MIMES:
            raise CampaignImportValidationError("CSV dosyasının MIME türü uyuşmuyor.")
        if data.startswith((b"PK\x03\x04", b"\xd0\xcf\x11\xe0")) or b"\x00" in data:
            raise CampaignImportValidationError("Geçerli bir CSV dosyası yükleyin.")
        return name, "text/csv"
    if suffix in {"xls", "xlsm"}:
        raise CampaignImportValidationError("Yalnız CSV veya makrosuz XLSX yükleyin.")
    if suffix != "xlsx" or (mime and mime != _XLSX_MIME):
        raise CampaignImportValidationError("Yalnız CSV veya XLSX yükleyin.")
    members = _zip_members(data)
    if "[content_types].xml" not in members or "xl/workbook.xml" not in members:
        raise CampaignImportValidationError("Geçerli bir XLSX dosyası yükleyin.")
    if any("vbaproject" in member for member in members) or {
        "encryptioninfo",
        "encryptedpackage",
    } & members:
        raise CampaignImportValidationError("Makrolu veya şifreli Excel dosyaları kabul edilmez.")
    return name, _XLSX_MIME


def build_object_key(tenant_id: UUID, import_id: UUID, sha256: str) -> str:
    """Do not place the original filename or raw phone data in storage keys."""

    return f"{_IMPORT_OBJECT_PREFIX}{tenant_id}/{import_id}/{sha256[:16]}.bin"


async def create_campaign_import(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    workflow_id: UUID,
    uploader_id: UUID,
    client_operation_id: UUID,
    country_code: str,
    filename: str,
    claimed_mime: str | None,
    data: bytes,
) -> tuple[CampaignImport, bool]:
    """Persist a private file and its idempotent import record.

    The caller must commit the surrounding workflow transaction before the
    delayed task can observe the row. ``queue_import(record)`` intentionally
    adds a short delay for this reason.
    """

    country = _normalise_country_code(country_code)
    safe_name, mime_type = validate_upload(filename, claimed_mime, data)
    sha256 = hashlib.sha256(data).hexdigest()
    existing = await db.scalar(
        select(CampaignImport).where(
            CampaignImport.tenant_id == tenant_id,
            CampaignImport.workflow_id == workflow_id,
            CampaignImport.client_operation_id == client_operation_id,
        )
    )
    if existing is not None:
        if existing.sha256 != sha256 or existing.country_code != country:
            raise CampaignImportConflictError(
                "Aynı işlem kimliği farklı dosya veya ülke seçimiyle kullanılamaz."
            )
        return existing, False

    import_id = uuid4()
    object_key = build_object_key(tenant_id, import_id, sha256)
    storage = get_campaign_import_storage()
    await storage.put_bytes(object_key, data, mime_type)
    row = CampaignImport(
        id=import_id,
        tenant_id=tenant_id,
        workflow_id=workflow_id,
        uploader_id=uploader_id,
        client_operation_id=client_operation_id,
        object_key=object_key,
        original_filename=safe_name,
        sha256=sha256,
        mime_type=mime_type,
        status=CampaignImportStatus.QUEUED.value,
        country_code=country,
        expires_at=datetime.now(UTC) + timedelta(days=IMPORT_RETENTION_DAYS),
    )
    db.add(row)
    try:
        await db.flush()
    except Exception:
        await storage.delete(object_key)
        raise
    return row, True


async def get_campaign_import(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    import_id: UUID,
    workflow_id: UUID | None = None,
    uploader_id: UUID | None = None,
    lock: bool = False,
) -> CampaignImport | None:
    stmt = select(CampaignImport).where(
        CampaignImport.id == import_id,
        CampaignImport.tenant_id == tenant_id,
    )
    if workflow_id is not None:
        stmt = stmt.where(CampaignImport.workflow_id == workflow_id)
    if uploader_id is not None:
        stmt = stmt.where(CampaignImport.uploader_id == uploader_id)
    if lock:
        stmt = stmt.with_for_update().execution_options(populate_existing=True)
    return cast(CampaignImport | None, await db.scalar(stmt))


async def select_phone_column(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    import_id: UUID,
    phone_column: str,
    workflow_id: UUID | None = None,
) -> CampaignImport:
    """Persist a manager-selected header and make the import parseable again."""

    row = await get_campaign_import(
        db,
        tenant_id=tenant_id,
        import_id=import_id,
        workflow_id=workflow_id,
        lock=True,
    )
    if row is None:
        raise CampaignImportValidationError("İçe aktarım bulunamadı.")
    if row.status != CampaignImportStatus.AWAITING_MAPPING.value:
        raise CampaignImportConflictError(
            "Telefon sütunu yalnız eşleşme belirsiz olduğunda seçilebilir."
        )
    selected = phone_column.strip()
    if selected not in row.columns:
        raise CampaignImportValidationError("Telefon sütunu dosyada bulunamadı.")
    row.phone_column = selected
    row.status = CampaignImportStatus.QUEUED.value
    # A retry always receives a new worker lease. This is also defensive for
    # rows recovered from a process that died after selecting a mapping.
    row.parse_token = None
    row.failure_reason = None
    row.summary = {}
    await db.flush()
    return row


async def record_consent(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    import_id: UUID,
    workflow_id: UUID,
    consent_source: str,
    consent_note: str | None,
) -> CampaignImport:
    """Persist the manager's attestation; audit logging remains with the workflow action."""

    row = await get_campaign_import(
        db,
        tenant_id=tenant_id,
        import_id=import_id,
        workflow_id=workflow_id,
        lock=True,
    )
    if row is None:
        raise CampaignImportValidationError("İçe aktarım bulunamadı.")
    source = consent_source.strip()
    note = (consent_note or "").strip() or None
    if source not in CONSENT_SOURCES:
        raise CampaignImportValidationError("İzin kaynağını belirtin.")
    if note is not None and len(note) > 2000:
        raise CampaignImportValidationError("İzin notu en fazla 2.000 karakter olabilir.")
    row.consent_source = source
    row.consent_note = note
    await db.flush()
    return row


def queue_import(
    campaign_import: CampaignImport | UUID,
    tenant_id: UUID | None = None,
    *,
    phone_column: str | None = None,
) -> None:
    """Enqueue parsing after the caller commits its transaction.

    A bare UUID has no safe way to recover its tenant under RLS, so callers
    must pass the ORM row or explicitly pass ``tenant_id``.
    """

    import_id: UUID
    tenant: UUID | None
    if isinstance(campaign_import, CampaignImport):
        import_id = campaign_import.id
        tenant = campaign_import.tenant_id
    else:
        import_id = campaign_import
        tenant = tenant_id
    if tenant is None:
        raise ValueError("tenant_id is required when queueing an import by id")
    from src.workers.campaign_imports import parse_campaign_import

    parse_campaign_import.apply_async(
        args=(str(tenant), str(import_id)),
        kwargs={"phone_column": phone_column} if phone_column else {},
        countdown=1,
    )


# Alias kept intentionally readable at workflow call sites.
queue_parse = queue_import


async def eligible_rows(
    db: AsyncSession, *, tenant_id: UUID, import_id: UUID
) -> list[CampaignImportRow]:
    imported = await get_campaign_import(db, tenant_id=tenant_id, import_id=import_id)
    if imported is None or imported.status != CampaignImportStatus.READY.value:
        raise CampaignImportConflictError("İçe aktarım gönderime hazır değil.")
    return list(
        (
            await db.scalars(
                select(CampaignImportRow)
                .where(
                    CampaignImportRow.tenant_id == tenant_id,
                    CampaignImportRow.campaign_import_id == import_id,
                    CampaignImportRow.status == CampaignImportRowStatus.ELIGIBLE.value,
                )
                .order_by(CampaignImportRow.row_number)
            )
        ).all()
    )


def public_import_view(row: CampaignImport) -> dict[str, Any]:
    """Return manager-UI aggregates and masked examples, never raw source cells.

    This shape deliberately still contains manager-facing metadata such as a
    filename and selectable header. The router applies a stricter status-only
    projection before any value crosses the LLM boundary.
    """

    examples = list((row.summary or {}).get("examples", []))
    return {
        "id": str(row.id),
        "import_id": str(row.id),
        "workflow_id": str(row.workflow_id),
        "filename": row.original_filename,
        "mime_type": row.mime_type,
        "status": row.status,
        "import_status": row.status,
        "country_code": row.country_code,
        "phone_column": row.phone_column,
        "columns": list(row.columns or []),
        "counts": {
            "total": row.total_rows,
            "eligible": row.eligible_count,
            "invalid": row.invalid_count,
            "duplicate": row.duplicate_count,
            "blocked": row.blocked_count,
        },
        "examples": examples,
        "errors": examples,
        "failure_reason": row.failure_reason,
        # This view is also copied into Workflow.state (JSONB). Keep every
        # value JSON-safe rather than relying on FastAPI's response encoder.
        "expires_at": row.expires_at.isoformat(),
    }


_EXPIRED_IMPORT_FAILURE = "İçe aktarma gizlilik saklama süresi doldu."
_IMPORT_WORKFLOW_FIELD_KEYS = frozenset(
    {
        "campaign_file",
        "phone_column",
        "country_code",
        "consent_source",
        "consent_note",
        "consent_confirmed",
    }
)


def expired_import_view() -> dict[str, str]:
    """A non-sensitive tombstone for cards after the 30-day cleanup."""

    return {"status": "failed", "failure_reason": _EXPIRED_IMPORT_FAILURE}


def redact_expired_import_fields(fields: dict[str, Any]) -> dict[str, Any]:
    """Remove file-derived workflow fields while preserving durable receipts."""

    return {
        key: value for key, value in fields.items() if key not in _IMPORT_WORKFLOW_FIELD_KEYS
    }


def redact_expired_import_request(value: Any) -> Any:
    """Strip file fields from an action receipt already tied to an expiry."""

    if isinstance(value, list):
        return [redact_expired_import_request(item) for item in value]
    if not isinstance(value, dict):
        return value
    return {
        key: (
            redact_expired_import_fields(item)
            if key == "fields" and isinstance(item, dict)
            else redact_expired_import_request(item)
        )
        for key, item in value.items()
    }


def _has_import_reference(value: dict[str, Any], import_id: UUID) -> bool:
    target = str(import_id)
    if str(value.get("import_id") or "") == target:
        return True
    for candidate in (
        value.get("campaign_import"),
        value.get("output", {}).get("campaign_import")
        if isinstance(value.get("output"), dict)
        else None,
    ):
        if isinstance(candidate, dict) and str(
            candidate.get("import_id") or candidate.get("id") or ""
        ) == target:
            return True
    return False


def redact_expired_import_artifacts(value: Any, import_id: UUID) -> Any:
    """Remove import-card replicas from durable workflow/chat JSON.

    The import table and rows are not the only persistence locations: workflow
    action receipts and chat-turn snapshots contain manager cards for
    idempotency. Once the source expires, those copies must not keep its
    filename, header names, masked examples or consent-note duplicate alive.
    The outbound receipt itself is intentionally not passed through here.
    """

    if isinstance(value, list):
        return [redact_expired_import_artifacts(item, import_id) for item in value]
    if not isinstance(value, dict):
        return value
    target = _has_import_reference(value, import_id)
    # A standalone public import view can appear directly in an idempotency
    # response/result rather than below ``campaign_import``. Replace the
    # entire view before recursion so filename, headers and masked examples do
    # not survive merely because there is no surrounding workflow wrapper.
    if target and any(
        key in value for key in {"filename", "mime_type", "columns", "examples", "counts"}
    ):
        return expired_import_view()
    redacted = {
        key: redact_expired_import_artifacts(item, import_id) for key, item in value.items()
    }
    if not target:
        return redacted
    if isinstance(value.get("campaign_import"), dict):
        redacted["campaign_import"] = expired_import_view()
    if "import_id" in value:
        redacted.pop("import_id", None)
        redacted["import_status"] = "expired"
    if isinstance(value.get("fields"), dict):
        redacted["fields"] = redact_expired_import_fields(value["fields"])
    return redacted


def _header_key(value: str) -> str:
    folded = unicodedata.normalize("NFKD", value.replace("ı", "i").replace("İ", "i"))
    return "".join(char for char in folded.casefold() if char.isalnum())


def _normalise_headers(values: tuple[Any, ...] | list[Any]) -> list[str]:
    if len(values) > _MAX_COLUMNS:
        raise CampaignImportValidationError("Dosyada çok fazla sütun var.")
    result: list[str] = []
    used: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        label = _cell_text(value).strip() or f"Sütun {index}"
        if len(label) > 255:
            raise CampaignImportValidationError("Sütun adı çok uzun.")
        seen = used.get(label, 0) + 1
        used[label] = seen
        result.append(label if seen == 1 else f"{label} ({seen})")
    if not result:
        raise CampaignImportValidationError("Dosyada başlık satırı bulunamadı.")
    return result


def _auto_phone_column(columns: list[str]) -> str | None:
    candidates = [column for column in columns if _header_key(column) in _PHONE_HEADER_KEYS]
    return candidates[0] if len(candidates) == 1 else None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_csv(data: bytes, phone_column: str | None) -> ParsedImport:
    decoded: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1254"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise CampaignImportValidationError("CSV UTF-8 veya Windows-1254 kodlamalı olmalı.")
    try:
        dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(decoded, newline=""), dialect)
    try:
        raw_headers = next(reader)
    except StopIteration as exc:
        raise CampaignImportValidationError("Dosyada başlık satırı bulunamadı.") from exc
    columns = _normalise_headers(raw_headers)
    selected = phone_column or _auto_phone_column(columns)
    if selected is not None and selected not in columns:
        raise CampaignImportValidationError("Seçilen telefon sütunu dosyada bulunamadı.")
    if selected is None:
        return ParsedImport(columns=columns, phone_column=None, phone_values=[])
    selected_index = columns.index(selected)
    return ParsedImport(
        columns=columns,
        phone_column=selected,
        phone_values=_phone_values_from_rows(reader, selected_index),
    )


def _parse_xlsx(data: bytes, phone_column: str | None) -> ParsedImport:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:  # pragma: no cover - deployment dependency
        raise CampaignImportValidationError("Excel ayrıştırıcısı kullanılabilir değil.") from exc
    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=True, keep_links=False
        )
    except Exception as exc:
        raise CampaignImportValidationError("Excel dosyası okunamadı.") from exc
    try:
        worksheet = next((sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"), None)
        if worksheet is None:
            raise CampaignImportValidationError("Dosyada görünür bir çalışma sayfası yok.")
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise CampaignImportValidationError("Dosyada başlık satırı bulunamadı.") from exc
        columns = _normalise_headers(raw_headers)
        selected = phone_column or _auto_phone_column(columns)
        if selected is not None and selected not in columns:
            raise CampaignImportValidationError("Seçilen telefon sütunu dosyada bulunamadı.")
        if selected is None:
            return ParsedImport(columns=columns, phone_column=None, phone_values=[])
        return ParsedImport(
            columns=columns,
            phone_column=selected,
            phone_values=_phone_values_from_rows(rows, columns.index(selected)),
        )
    finally:
        workbook.close()


def _phone_values_from_rows(rows: Any, selected_index: int) -> list[tuple[int, str]]:
    values: list[tuple[int, str]] = []
    for row_number, source in enumerate(rows, start=2):
        fields = [_cell_text(value) for value in source]
        if not any(fields):
            continue
        if len(values) >= MAX_IMPORT_ROWS:
            raise CampaignImportValidationError("Dosyada en fazla 10.000 veri satırı olabilir.")
        values.append((row_number, fields[selected_index] if selected_index < len(fields) else ""))
    return values


def parse_import_bytes(
    data: bytes, mime_type: str, phone_column: str | None = None
) -> ParsedImport:
    """Read a bounded CSV/XLSX payload without retaining unrelated cell data."""

    if len(data) > MAX_IMPORT_BYTES:
        raise CampaignImportValidationError("Dosya en fazla 10 MiB olabilir.")
    if mime_type == "text/csv":
        return _parse_csv(data, phone_column)
    if mime_type == _XLSX_MIME:
        return _parse_xlsx(data, phone_column)
    raise CampaignImportValidationError("Desteklenmeyen içe aktarma türü.")


def normalise_phone(raw: str, country_code: str) -> str | None:
    value = raw.strip()
    if not value or len(value) > _MAX_PHONE_CELL_CHARS:
        return None
    try:
        parsed = phonenumbers.parse(value, country_code)
    except phonenumbers.NumberParseException:
        return None
    if not phonenumbers.is_valid_number(parsed):
        return None
    return phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164)


def mask_phone(value: str | None) -> str | None:
    if not value:
        return None
    # Never derive a display value from arbitrary invalid source text. A CSV
    # cell can contain a name, note, or another identifier; only digits are
    # safe to retain in a masked phone preview.
    if len(value) > _MAX_PHONE_CELL_CHARS:
        return "••••"
    digits = re.sub(r"\D", "", value)
    if len(digits) <= 4:
        return "••••"
    prefix = "+" if value.strip().startswith("+") else ""
    return f"{prefix}{digits[:2]}••••{digits[-3:]}"


async def known_block_reasons(
    db: AsyncSession, tenant_id: UUID, phones: set[str]
) -> dict[str, str]:
    """Return only hard blocks already known locally.

    Consent evidence and the rolling outbound cooldown intentionally remain a
    final queue/worker check, because a manager may add consent evidence after
    parsing. Opt-outs, invalid known contacts, and blacklisted leads cannot be
    overridden by that evidence and are safe to surface early.
    """

    if not phones:
        return {}
    from src.modules.compliance.models import ComplianceCheck, ComplianceResult, OptOut
    from src.modules.discovery.models import (
        ConsentStatus,
        ContactType,
        Lead,
        LeadContact,
        LeadStatus,
    )

    reasons: dict[str, str] = {}
    values = list(phones)
    for start in range(0, len(values), 1000):
        chunk = values[start : start + 1000]
        opted_out = await db.scalars(
            select(OptOut.phone_e164).where(
                OptOut.tenant_id == tenant_id,
                OptOut.phone_e164.in_(chunk),
            )
        )
        for phone in opted_out:
            reasons[phone] = "Müşteri iletişimi durdurmuş."
        contacts = await db.execute(
            select(
                LeadContact.id,
                LeadContact.normalized_value,
                LeadContact.consent_status,
                LeadContact.is_valid,
                Lead.status,
            )
            .join(Lead, Lead.id == LeadContact.lead_id)
            .where(
                LeadContact.tenant_id == tenant_id,
                LeadContact.type == ContactType.PHONE,
                LeadContact.normalized_value.in_(chunk),
            )
        )
        contact_phones: dict[UUID, str] = {}
        for contact_id, phone, consent_status, is_valid, lead_status in contacts:
            contact_phones[contact_id] = phone
            if phone in reasons:
                continue
            if consent_status == ConsentStatus.OPT_OUT or not is_valid:
                reasons[phone] = "Numara geçersiz veya müşteri iletişimi durdurmuş."
            elif lead_status in {LeadStatus.BLACKLISTED, LeadStatus.BLOCKED_BY_COMPLIANCE}:
                reasons[phone] = "Müşteri şirketin engelli listesinde."
        if contact_phones:
            checks = await db.execute(
                select(
                    ComplianceCheck.contact_id,
                    ComplianceCheck.result,
                    ComplianceCheck.next_allowed_at,
                )
                .where(
                    ComplianceCheck.tenant_id == tenant_id,
                    ComplianceCheck.contact_id.in_(list(contact_phones)),
                )
                .order_by(ComplianceCheck.contact_id, ComplianceCheck.created_at.desc())
            )
            seen_checks: set[UUID] = set()
            now = datetime.now(UTC)
            for contact_id, result, next_allowed_at in checks:
                if contact_id is None or contact_id in seen_checks:
                    continue
                seen_checks.add(contact_id)
                checked_phone = contact_phones.get(contact_id)
                if not checked_phone or checked_phone in reasons:
                    continue
                if result == ComplianceResult.BLOCK:
                    reasons[checked_phone] = "Uyumluluk kontrolü bu alıcıyı engelliyor."
                elif result == ComplianceResult.DEFER and (
                    next_allowed_at is None or next_allowed_at > now
                ):
                    reasons[checked_phone] = "Uyumluluk kontrolü bu alıcı için henüz bekleme istiyor."
    return reasons


async def build_campaign_import_rows(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    campaign_import_id: UUID,
    country_code: str,
    phone_values: list[tuple[int, str]],
) -> tuple[list[CampaignImportRow], dict[str, int], list[dict[str, Any]]]:
    """Normalize, de-duplicate and classify a bounded import without raw data.

    The returned row records contain only an E.164 number when valid and a
    masked display value. Raw source cell text is deliberately never persisted.
    """

    normalized = [normalise_phone(value, country_code) for _, value in phone_values]
    blockers = await known_block_reasons(db, tenant_id, {phone for phone in normalized if phone})
    seen: set[str] = set()
    rows: list[CampaignImportRow] = []
    counts = {
        CampaignImportRowStatus.ELIGIBLE.value: 0,
        CampaignImportRowStatus.INVALID.value: 0,
        CampaignImportRowStatus.DUPLICATE.value: 0,
        CampaignImportRowStatus.BLOCKED.value: 0,
    }
    examples: list[dict[str, Any]] = []
    for (row_number, raw), phone in zip(phone_values, normalized, strict=True):
        status: CampaignImportRowStatus
        reason: str | None
        if phone is None:
            status = CampaignImportRowStatus.INVALID
            reason = "Geçerli telefon numarası değil."
        elif phone in seen:
            status = CampaignImportRowStatus.DUPLICATE
            reason = "Dosyada tekrar eden numara."
        else:
            seen.add(phone)
            reason = blockers.get(phone)
            status = CampaignImportRowStatus.BLOCKED if reason else CampaignImportRowStatus.ELIGIBLE
        counts[status.value] += 1
        masked = mask_phone(phone or raw)
        rows.append(
            CampaignImportRow(
                tenant_id=tenant_id,
                campaign_import_id=campaign_import_id,
                row_number=row_number,
                phone_e164=phone,
                masked_phone=masked,
                status=status.value,
                reason=reason,
            )
        )
        if status != CampaignImportRowStatus.ELIGIBLE and len(examples) < 10:
            examples.append(
                {
                    "row_number": row_number,
                    "phone": masked,
                    "status": status.value,
                    "reason": reason,
                }
            )
    return rows, counts, examples
