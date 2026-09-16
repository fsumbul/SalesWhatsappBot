"""Local text extraction for uploaded documents and fetched web pages.

Everything here runs in-process (pypdf, openpyxl, trafilatura, BeautifulSoup):
raw tenant files and pages never leave the deployment. The output is a list of
``TextUnit`` objects with a stable locator (``file.pdf#page=3``,
``file.xlsx#sheet=Fiyat&rows=2-41``, a page URL) that candidates cite as
evidence.
"""

from __future__ import annotations

import csv
import io
import json
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any
from urllib.parse import urljoin, urlparse

from bs4 import BeautifulSoup

from .chunking import normalize_whitespace

EXTRACTOR_VERSION = "2026.09.2"

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_SUPPORTED_MIMES = {
    "application/pdf",
    _XLSX_MIME,
    "text/csv",
    "text/markdown",
    "text/plain",
    "text/html",
}
_EXTENSION_MIME = {
    ".pdf": "application/pdf",
    ".xlsx": _XLSX_MIME,
    ".csv": "text/csv",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
    ".html": "text/html",
    ".htm": "text/html",
}
_ROWS_PER_UNIT = 40
_IMAGE_URL_SKIP_RE = re.compile(
    r"(logo|icon|sprite|favicon|banner|placeholder|loading|spinner|avatar|flag|badge|"
    r"pixel|tracking|\.svg(\?|$)|\.gif(\?|$))",
    re.IGNORECASE,
)


class UnsupportedDocumentError(ValueError):
    """The uploaded file type cannot be extracted locally."""


@dataclass(frozen=True)
class TextUnit:
    locator: str
    text: str
    title: str | None = None
    meta: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ImageRef:
    url: str
    alt: str = ""
    context: str = ""
    width: int | None = None
    height: int | None = None
    source: str = "img"  # img | og | jsonld


@dataclass(frozen=True)
class PageExtract:
    url: str
    title: str | None
    text: str
    links: tuple[str, ...]
    images: tuple[ImageRef, ...]
    product: dict[str, Any]
    canonical: str | None = None


def sniff_mime(filename: str, declared: str | None, data: bytes) -> str:
    """Trust magic bytes first, then the extension; reject anything else."""

    if data.startswith(_PDF_MAGIC):
        return "application/pdf"
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if data.startswith(_ZIP_MAGIC) and extension == ".xlsx":
        return _XLSX_MIME
    if data.startswith(_ZIP_MAGIC):
        raise UnsupportedDocumentError("Yalnızca .xlsx tablo dosyaları desteklenir.")
    by_extension = _EXTENSION_MIME.get(extension)
    declared_clean = (declared or "").split(";")[0].strip().lower()
    mime = by_extension or (declared_clean if declared_clean in _SUPPORTED_MIMES else "")
    if mime not in _SUPPORTED_MIMES or mime in {"application/pdf", _XLSX_MIME}:
        raise UnsupportedDocumentError(
            "Desteklenen türler: PDF, XLSX, CSV, Markdown, düz metin, HTML."
        )
    try:
        data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedDocumentError("Metin dosyası UTF-8 olmalıdır.") from exc
    return mime


def extract_document(filename: str, mime_type: str, data: bytes) -> list[TextUnit]:
    if mime_type == "application/pdf":
        return _extract_pdf(filename, data)
    if mime_type == _XLSX_MIME:
        return _extract_xlsx(filename, data)
    if mime_type == "text/csv":
        return _extract_csv(filename, data.decode("utf-8", errors="replace"))
    if mime_type == "text/html":
        page = html_to_page(f"file://{filename}", data.decode("utf-8", errors="replace"))
        return [TextUnit(locator=filename, text=page.text, title=page.title)] if page.text else []
    if mime_type in {"text/markdown", "text/plain"}:
        return _extract_markdown(filename, data.decode("utf-8", errors="replace"))
    raise UnsupportedDocumentError(f"Unsupported mime type {mime_type}")


def _extract_pdf(filename: str, data: bytes) -> list[TextUnit]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    units: list[TextUnit] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = normalize_whitespace(page.extract_text() or "")
        except Exception:  # a broken page must not lose the whole document
            text = ""
        if text:
            units.append(
                TextUnit(locator=f"{filename}#page={number}", text=text, meta={"page": number})
            )
    return units


def table_rows_to_units(
    filename: str,
    sheet: str,
    rows: list[list[str]],
    *,
    locator_for: Callable[[int, int], str] | None = None,
    extra_meta: dict[str, Any] | None = None,
) -> list[TextUnit]:
    """Render a header + rows grid as ``Tablo:`` units (XLSX, CSV and OCR tables alike)."""

    if not rows:
        return []
    header = [cell.strip() for cell in rows[0]]
    body = rows[1:] if any(header) else rows
    header = header if any(header) else [f"sütun{i + 1}" for i in range(len(rows[0]))]
    units: list[TextUnit] = []
    for start in range(0, len(body), _ROWS_PER_UNIT):
        block = body[start : start + _ROWS_PER_UNIT]
        lines = []
        for row in block:
            cells = [
                f"{header[i] if i < len(header) else f'sütun{i + 1}'}: {value.strip()}"
                for i, value in enumerate(row)
                if value and value.strip()
            ]
            if cells:
                lines.append(" | ".join(cells))
        if not lines:
            continue
        first, last = start + 2, start + 1 + len(block)
        locator = (
            locator_for(first, last)
            if locator_for is not None
            else f"{filename}#sheet={sheet}&rows={first}-{last}"
        )
        units.append(
            TextUnit(
                locator=locator,
                text=f"Tablo: {sheet}\n" + "\n".join(lines),
                title=sheet,
                meta={"sheet": sheet, "rows": [first, last], **(extra_meta or {})},
            )
        )
    return units


_table_rows_to_units = table_rows_to_units


def _extract_xlsx(filename: str, data: bytes) -> list[TextUnit]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    units: list[TextUnit] = []
    for sheet in workbook.worksheets:
        rows = [
            ["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)
        ]
        units.extend(_table_rows_to_units(filename, sheet.title, rows))
    return units


def _extract_csv(filename: str, text: str) -> list[TextUnit]:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    return _table_rows_to_units(filename, "csv", rows)


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$", re.MULTILINE)


def _extract_markdown(filename: str, text: str) -> list[TextUnit]:
    """Split on headings so locators point at the section that was cited."""

    matches = list(_HEADING_RE.finditer(text))
    if not matches:
        body = normalize_whitespace(text)
        return [TextUnit(locator=filename, text=body)] if body else []
    units: list[TextUnit] = []
    preface = normalize_whitespace(text[: matches[0].start()])
    if preface:
        units.append(TextUnit(locator=filename, text=preface))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        title = match.group(2).strip()
        body = normalize_whitespace(text[match.end() : end])
        if body:
            units.append(
                TextUnit(
                    locator=f"{filename}#{_slug(title)}",
                    text=f"{title}\n{body}",
                    title=title,
                )
            )
    return units


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")[:80] or "section"


# --- HTML pages --------------------------------------------------------------


def _absolute(base: str, href: str | None) -> str | None:
    if not href:
        return None
    href = href.strip()
    if href.startswith(("javascript:", "mailto:", "tel:", "data:", "#")):
        return None
    absolute = urljoin(base, href)
    parsed = urlparse(absolute)
    if parsed.scheme not in {"http", "https"}:
        return None
    return absolute.split("#", 1)[0]


def _int_attr(value: Any) -> int | None:
    try:
        return int(str(value).strip().rstrip("px")) if value else None
    except ValueError:
        return None


def _jsonld_products(soup: BeautifulSoup) -> list[dict[str, Any]]:
    products: list[dict[str, Any]] = []
    for script in soup.find_all("script", attrs={"type": "application/ld+json"}):
        try:
            payload = json.loads(script.string or "")
        except (TypeError, ValueError):
            continue
        stack: list[Any] = [payload]
        while stack:
            item = stack.pop()
            if isinstance(item, list):
                stack.extend(item)
            elif isinstance(item, dict):
                if "@graph" in item:
                    stack.append(item["@graph"])
                kind = item.get("@type")
                kinds = kind if isinstance(kind, list) else [kind]
                if any(str(k).lower() == "product" for k in kinds if k):
                    products.append(item)
    return products


def html_to_page(url: str, html: str) -> PageExtract:
    """Main text (trafilatura) plus structure (links, images, JSON-LD product)."""

    import trafilatura

    soup = BeautifulSoup(html, "html.parser")
    title_tag = soup.find("title")
    title = normalize_whitespace(title_tag.get_text(" ", strip=True))[:300] if title_tag else None
    og_title = soup.find("meta", attrs={"property": "og:title"})
    if og_title and og_title.get("content"):
        title = normalize_whitespace(str(og_title["content"]))[:300]

    text = (
        trafilatura.extract(
            html,
            url=url,
            include_tables=True,
            include_links=False,
            include_comments=False,
            favor_recall=True,
        )
        or ""
    )
    text = normalize_whitespace(text)
    if len(text) < 80:
        # Fallback: strip navigation/script noise and keep the body text.
        for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "form"]):
            tag.decompose()
        body = soup.find("body") or soup
        text = normalize_whitespace(body.get_text("\n", strip=True))

    links: list[str] = []
    for anchor in soup.find_all("a", href=True):
        absolute = _absolute(url, str(anchor.get("href")))
        if absolute and absolute not in links:
            links.append(absolute)

    images: list[ImageRef] = []
    seen: set[str] = set()

    def _add(ref: ImageRef) -> None:
        if ref.url in seen or _IMAGE_URL_SKIP_RE.search(ref.url):
            return
        seen.add(ref.url)
        images.append(ref)

    products = _jsonld_products(soup)
    product: dict[str, Any] = {}
    if products:
        first = products[0]
        product = {
            "name": str(first.get("name") or "")[:200],
            "description": normalize_whitespace(str(first.get("description") or ""))[:1000],
            "sku": str(first.get("sku") or "")[:80],
            "brand": (
                str(
                    first["brand"].get("name")
                    if isinstance(first.get("brand"), dict)
                    else first.get("brand") or ""
                )
            )[:120],
        }
        raw_images = first.get("image") or []
        for item in raw_images if isinstance(raw_images, list) else [raw_images]:
            candidate = item.get("url") if isinstance(item, dict) else item
            absolute = _absolute(url, str(candidate) if candidate else None)
            if absolute:
                _add(
                    ImageRef(
                        url=absolute, alt=product["name"], context=product["name"], source="jsonld"
                    )
                )
    og_type = soup.find("meta", attrs={"property": "og:type"})
    if og_type and "product" in str(og_type.get("content", "")).lower():
        product.setdefault("name", title or "")
        product["og_type"] = "product"
    og_image = soup.find("meta", attrs={"property": "og:image"})
    if og_image and og_image.get("content"):
        absolute = _absolute(url, str(og_image["content"]))
        if absolute:
            _add(ImageRef(url=absolute, alt=title or "", context=title or "", source="og"))

    for img in soup.find_all("img"):
        src = img.get("data-src") or img.get("src")
        srcset = img.get("srcset") or img.get("data-srcset")
        if srcset and not src:
            src = str(srcset).split(",")[-1].strip().split(" ")[0]
        absolute = _absolute(url, str(src) if src else None)
        if not absolute:
            continue
        heading = img.find_previous(["h1", "h2", "h3", "figcaption"])
        context = normalize_whitespace(heading.get_text(" ", strip=True))[:200] if heading else ""
        _add(
            ImageRef(
                url=absolute,
                alt=normalize_whitespace(str(img.get("alt") or ""))[:200],
                context=context,
                width=_int_attr(img.get("width")),
                height=_int_attr(img.get("height")),
            )
        )

    canonical_tag = soup.find("link", attrs={"rel": "canonical"})
    canonical = (
        _absolute(url, str(canonical_tag.get("href")))
        if canonical_tag and canonical_tag.get("href")
        else None
    )
    return PageExtract(
        url=url,
        title=title,
        text=text,
        links=tuple(links),
        images=tuple(images),
        product=product,
        canonical=canonical,
    )
